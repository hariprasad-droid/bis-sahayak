from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    subtitle.text = subtitle_text
    for p in subtitle.text_frame.paragraphs:
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(89, 89, 89)

def add_content_slide(prs, title_text, bullet_points):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = bullet_points[0]
    
    for point in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = point
        if point.startswith(" -") or point.startswith("-"):
            p.level = 1
        elif point.startswith("   -"):
            p.level = 2

# Slide 1 (Speaker 1)
add_title_slide(
    prs, 
    "BIS Sahayak", 
    "AI-Powered Intelligent Assistant for Indian Standards\n\nSIH26107 | Smart Automation\nTeam Rockstar (095)"
)

# Slide 2 (Speaker 2)
add_content_slide(
    prs, 
    "Problem & Solution", 
    [
        "The Problem:",
        " - Information on standards is scattered across thousands of PDFs and portals.",
        " - Manual searching wastes time and allows counterfeit products into the market.",
        "The Solution: BIS Sahayak",
        " - A multilingual, RAG-based conversational AI agent.",
        " - No guessing or hallucinations: retrieves exact clauses before answering.",
        " - Guides users step-by-step from product query to certification."
    ]
)

# Slide 3 (Speaker 3)
add_content_slide(
    prs,
    "Technical Approach",
    [
        "6-Step Methodology:",
        " - Query Understanding → Semantic Retrieval → Standard Mapping → Explainable Response → Multilingual Delivery → Human Escalation",
        "Technology Stack:",
        " - Python, FastAPI, React.js",
        " - LangChain (LLM Orchestration) & FAISS (Semantic Search)",
        " - Whisper (Speech Recognition for regional languages)",
        "Integration Ready:",
        " - API-ready for BIS CARE portal and UMANG."
    ]
)

# Slide 4 (Speaker 4)
add_content_slide(
    prs,
    "Feasibility & Viability",
    [
        "Technical Feasibility:",
        " - Built on proven, open-source RAG components. Microservice architecture for scalability.",
        " - Requires no user training (simple chat/voice UI).",
        "Market & Viability:",
        " - Over 6.3 crore MSMEs face this problem daily.",
        " - Reduces consulting costs and BIS staff workload.",
        "Sustainability:",
        " - Hybrid revenue model: Public assistant + Premium Enterprise API."
    ]
)

# Slide 5 (Speaker 5)
add_content_slide(
    prs,
    "Impact & Benefits",
    [
        "Measurable Impact (Before vs After):",
        " - Finding applicable standards drops from 5 to 1 (effort scale).",
        " - Certification guidance drops from 5 to 2.",
        "Broad Benefits:",
        " - Social: Builds trust in certified products.",
        " - Economic: Saves MSMEs time and consulting costs.",
        " - Operational: Frees BIS staff from repetitive queries.",
        " - Technological: Advances e-governance with NLP and multilingual AI."
    ]
)

# Slide 6 (Speaker 6)
add_content_slide(
    prs,
    "Research & Urgency",
    [
        "The Urgent Need:",
        " - Recent seizures of fake ISI-marked cables, appliances, and footwear (worth lakhs) highlight the need for instant verification.",
        "Primary Research Grounding:",
        " - Based directly on Indian Standards documents (e.g., IS 16102).",
        " - Adapted from proven regulatory RAG systems (LawPal, Pharma compliance).",
        "Summary:",
        " - BIS Sahayak reduces counterfeits, saves time, and eases the burden on officials."
    ]
)

# Save
output_path = r"c:\bis\sihh\BIS_Sahayak_6_Speaker_Pitch.pptx"
prs.save(output_path)
print(f"Presentation saved successfully to {output_path}")
